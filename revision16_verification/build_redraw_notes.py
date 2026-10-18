"""Create the standalone instruction-16 delivery note; never edit the manuscript."""
from pathlib import Path
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT

OUT=Path('/workspace/output_final/第二轮图片重绘说明.docx')


def font(run,size=10.5,bold=False):
    run.font.name='Times New Roman';run.font.size=Pt(size);run.bold=bold
    run.font.color.rgb=RGBColor.from_string('222222')
    fonts=run._element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn('w:eastAsia'),'宋体')
    return run


def paragraph(doc,text,bold=False,after=5,size=10.5):
    p=doc.add_paragraph();font(p.add_run(text),size,bold)
    f=p.paragraph_format;f.space_after=Pt(after);f.line_spacing=1.12
    return p


def heading(doc,text):
    p=paragraph(doc,text,True,4,11)
    p.paragraph_format.space_before=Pt(6);p.paragraph_format.keep_with_next=True
    return p


def border(cell,top=None,bottom=None):
    props=cell._tc.get_or_add_tcPr();b=props.find(qn('w:tcBorders'))
    if b is None:b=OxmlElement('w:tcBorders');props.append(b)
    for edge in ['top','left','bottom','right']:
        element=OxmlElement('w:'+edge)
        value=top if edge=='top' else bottom if edge=='bottom' else None
        element.set(qn('w:val'),'single' if value else 'nil')
        if value:element.set(qn('w:sz'),str(value));element.set(qn('w:color'),'434C54')
        b.append(element)


def main():
    doc=Document();s=doc.sections[0]
    s.page_width=Cm(21);s.page_height=Cm(29.7)
    s.top_margin=Cm(1.8);s.bottom_margin=Cm(1.7)
    s.left_margin=Cm(2.0);s.right_margin=Cm(2.0)
    normal=doc.styles['Normal'];normal.font.name='Times New Roman';normal.font.size=Pt(10.5)
    normal.element.rPr.rFonts.set(qn('w:eastAsia'),'宋体')
    doc.core_properties.title='第二轮图片重绘说明'
    doc.core_properties.subject='修改指令16：图1—图7正式重绘与科学内容保留说明'
    doc.core_properties.author=''
    p=paragraph(doc,'第二轮图片重绘说明',True,9,17)
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    paragraph(doc,'完成范围：图1—图7全部重绘。以第15轮图片优化后的最新Word为底稿，仅替换正文图片；图号、图题和正文引用保持不变，保留作者已经修正的题目、摘要格式、页眉页脚和最后一页分页处理。')
    paragraph(doc,'数据方式：全部基于原始观测CSV或已保存的预测、规划实验输出精确重绘。未使用截图估值、图像读数或缺失数据猜测；未重跑训练、排程优化或新增实验。',after=6)
    heading(doc,'逐图重绘内容')
    rows=[
        ('图1','建立“年度覆盖总览＋逐单元观测细节”层级；矩阵改用规整单元格，强化有观测／缺测区分，统一年份列与图例。','原始观测及年度覆盖CSV'),
        ('图2','统一三模型散点、对角参考线和误差标注；误差比较改为水平分组柱，增加直接数值与灰度纹理，保留模型顺序。','已保存预测与误差比较CSV'),
        ('图3','以空心点和醒目叉号区分53个估计初值、6个当年实测；加强零差线与差值正负区分，保留原直方分箱和频数。','已保存2024年初始状态CSV'),
        ('图4','重设计箱线、面板标题及统计图例；增加已有均值菱形，保留中位数、1.5×IQR须线和离群圆点。首次年差最低档n=28，其余各档n=30。','已保存配对运行及汇总CSV'),
        ('图5','以模型与统计类型分层编码：可行均值±样本SD、独立最好值及既有可行目标下界；标明LOCF有限搜索不足，保留原局部下降。','已保存资源响应及下界CSV'),
        ('图6','将平均评分、最低评分和年度养护数组织成一致三联图；统一年份网格，清楚呈现75／60分阈值、年度能力和各年次数，保留零次数。','已保存代表方案年度CSV'),
        ('图7','统一三种d情景的颜色、线型和点形；重组均值／样本SD／最好值图例，增加原值标注与首次养护年份点列；仅作横向排布，年份原值不变。','已保存敏感性汇总、年度及首次养护CSV'),
    ]
    table=doc.add_table(rows=1,cols=3);table.alignment=WD_TABLE_ALIGNMENT.CENTER
    table.autofit=False;widths=[1.2,11.0,4.8]
    for column,width in zip(table.columns,widths):column.width=Cm(width)
    for cell,width,text in zip(table.rows[0].cells,widths,['图号','主要视觉升级','数据依据']):
        cell.width=Cm(width);font(cell.paragraphs[0].add_run(text),9.5,True)
        border(cell,top=8,bottom=4)
    for index,row in enumerate(rows):
        cells=table.add_row().cells
        for cell,width,text in zip(cells,widths,row):
            cell.width=Cm(width);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(3);p.paragraph_format.space_before=Pt(3)
            p.paragraph_format.line_spacing=1.06;font(p.add_run(text),9.2)
            border(cell,bottom=8 if index==len(rows)-1 else None)
        trPr=table.rows[-1]._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');trPr.append(cant)
    heading(doc,'统计语义与保留边界')
    paragraph(doc,'图4新增均值直接读取paired_schedule_comparison_summary_15y.csv；图5新增目标下界直接读取resource_lower_bound_summary_15y.csv。均为项目中已有结果，不是新实验。已验证可行下界单独展示，不计入本档30次独立运行的均值、样本SD或最好值统计；NA未置为0。',after=4,size=10)
    paragraph(doc,'科研数值改动：否。正文科研结论改动：否。未平滑曲线、修改趋势、删除不利结果或改变样本数量；原数据坐标、误差棒、箱线结构和柱形数值均按冻结输入核对。',after=4,size=10)
    heading(doc,'交付与检查')
    paragraph(doc,'最终论文：小论文_中文终稿候选_高质量图片重绘版.docx，15页。单图输出至output_figures_redrawn/：图1—图7各含600dpi PNG、SVG及PDF，共21个文件；保留原Word图框的物理宽高，避免拉伸变形。',after=4,size=10)
    paragraph(doc,'output_scripts/提供统一重绘入口、共享样式和分图脚本，支持中文／英文切换；随附22份CSV，可独立读取复用。核验记录包括原始数组与图形数据比对、源文件哈希和文字边界检查。',after=4,size=10)
    paragraph(doc,'全部15页逐页核验通过，文字位置、原图框及图片以外的版面保持一致。未发现遗留裁字、图形变形或科研数据缺失；图号与正文引用一致，文字、数值、表格、参考文献及作者人工修正均予保留。当前无未完成项。',after=0,size=10)
    # Keep the standalone note compact; do not introduce a trailing page break.
    footer=s.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    OUT.parent.mkdir(exist_ok=True,parents=True);doc.save(OUT)
    print(OUT)


if __name__=='__main__':main()
