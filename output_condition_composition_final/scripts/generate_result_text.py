#!/usr/bin/env python3
"""Create the standalone Chinese caption and manuscript text from saved data.

No manuscript, research dataset, optimizer, or existing figure is modified.
Usage: python scripts/generate_result_text.py
The packaged input_data directory is preferred; --source-dir overrides it.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import re
from pathlib import Path

import numpy as np
import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


HERE = Path(__file__).resolve().parent.parent
ORDER = ["no_maintenance", "low_resource", "medium_resource", "high_resource"]
EXPECTED_COUNTS = [[0, 54, 5, 0, 0], [19, 39, 1, 0, 0],
                   [40, 18, 1, 0, 0], [58, 1, 0, 0, 0]]
EXPECTED_ACTIONS = [0, 58, 93, 114]


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def character_counts(text: str) -> dict:
    return {"cjk_han_characters": len(re.findall(r"[\u4e00-\u9fff]", text)),
            "non_whitespace_characters": len(re.sub(r"\s", "", text))}


def read_and_check(source_dir: Path):
    composition_path = source_dir / "condition_composition_by_year.csv"
    scores_path = source_dir / "unit_year_scores.csv"
    manifest_path = source_dir / "scenario_selection.json"
    composition = pd.read_csv(composition_path, float_precision="round_trip")
    scores = pd.read_csv(scores_path, float_precision="round_trip")
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    scenarios = {item["scenario"]: item for item in manifest["scenarios"]}
    years = list(range(2025, 2040))
    assert manifest["years"] == years
    assert manifest["n_units"] == 59 and manifest["initial_state_year"] == 2024
    assert manifest["initial_condition"] == "LOCF"
    assert manifest["d"] == 0.65 and manifest["R"] == 6
    assert len(composition) == 4 * 15 * 5 and len(scores) == 4 * 15 * 59
    assert set(composition.scenario) == set(ORDER)
    assert set(scores.scenario) == set(ORDER)
    assert not composition.duplicated(["scenario", "year", "grade"]).any()
    assert not scores.duplicated(["scenario", "unit_id", "year"]).any()

    observed = []
    for name, expected_counts, expected_actions in zip(ORDER, EXPECTED_COUNTS, EXPECTED_ACTIONS):
        rows = composition.loc[composition.scenario == name]
        unit_rows = scores.loc[scores.scenario == name]
        meta = scenarios[name]
        assert sorted(rows.year.unique().tolist()) == years
        assert sorted(unit_rows.year.unique().tolist()) == years
        assert rows.initial_condition.eq("LOCF").all()
        assert rows.d.eq(0.65).all() and rows.R.eq(6).all()
        assert rows.B.eq(meta["B"]).all()
        assert meta["initial_condition"] == "LOCF" and meta["d"] == .65 and meta["R"] == 6
        assert meta["start_year"] == 2025 and meta["end_year"] == 2039
        assert meta["planning_horizon"] == 15
        counts_by_year = rows.groupby("year")["count"].sum()
        percentages_by_year = rows.groupby("year")["percentage"].sum()
        assert counts_by_year.eq(59).all()
        assert np.allclose(percentages_by_year, 100., rtol=0., atol=1e-10)
        assert rows.loc[rows.grade.isin([4, 5]), "count"].eq(0).all()

        annual = []
        for year in years:
            c = rows.loc[rows.year == year].set_index("grade").sort_index()
            u = unit_rows.loc[unit_rows.year == year]
            assert len(u) == 59 and u.unit_id.nunique() == 59
            s = u.score.to_numpy()
            assert np.isfinite(s).all() and ((s >= 0) & (s <= 100)).all()
            # Exact interval tests: boundaries stay in the lower score class.
            g = np.select([s > 90, s > 75, s > 60, s > 40], [1, 2, 3, 4], default=5)
            counts = np.bincount(g, minlength=6)[1:6]
            assert np.array_equal(c.index.to_numpy(), np.arange(1, 6))
            assert np.array_equal(c["count"].to_numpy(), counts)
            assert np.allclose(c["percentage"].to_numpy(), counts / 59 * 100,
                               rtol=0., atol=1e-10)
            action_count = int(u.action.sum())
            assert u.action.isin([0, 1]).all()
            assert c.maintenance_actions.eq(action_count).all()
            annual.append(action_count)

        counts_2039 = rows.loc[rows.year == 2039].sort_values("grade")["count"].astype(int).tolist()
        assert counts_2039 == expected_counts
        assert annual == meta["annual_actions"]
        assert sum(annual) == expected_actions == meta["total_actions"]
        observed.append({"scenario": name, "B": meta["B"], "counts_2039": counts_2039,
                         "annual_actions": annual, "total_actions": sum(annual)})

    baseline = scores.loc[scores.scenario == "no_maintenance"]
    failing_years = [int(year) for year, group in baseline.groupby("year")
                     if (group.score <= 60).any() or int((group.score > 75).sum()) < 57]
    assert min(failing_years) == 2036
    assert not scenarios["no_maintenance"]["feasible_management_strategy"]
    assert all(scenarios[name]["feasible_management_strategy"] for name in ORDER[1:])
    hashes = {path.name: sha256(path) for path in [composition_path, scores_path, manifest_path]}
    return observed, hashes


def make_content(observed):
    # Every number in the result paragraph is read from the saved composition CSV.
    counts = ["、".join(str(i) for i in item["counts_2039"][:3]) for item in observed]
    B = [item["B"] for item in observed]
    title = "不同养护资源情景下桥梁检测单元评分等级构成的年度变化"
    caption = (
        "59个桥梁检测单元在2025—2039年的年末评分区间构成。四情景同用2024年LOCF初值、"
        "d=0.65分/年、R=6分及相同的基础递推与管理条件。"
        "（a）为固定零行动的无养护比较基线（B=0）；（b）—（d）依次为低、中、高资源养护情景，"
        f"B分别显示为{B[1]:.3f}、{B[2]:.3f}、{B[3]:.3f}，计算保留原始精度。"
        "B为以总面积归一化的累计养护面积份额上限，非货币预算。"
        "各类比例以59个单元为分母，按养护后年末连续模拟评分划分，非完整规范检测的正式技术状况评定。"
        "所有模拟情景中均未出现4类和5类评分区间的检测单元。"
        "下方条带表示当年执行养护的检测单元数，跨年行动分别计次。"
        "无养护基线自2036年起不满足管理约束，不作为可行养护方案。"
    )
    result = (
        "基于年度模拟评分，比较无养护及低、中、高资源养护情景下59个检测单元的评分区间构成。"
        f"2039年，无养护情景的1—3类数量依次为{counts[0]}个；低资源情景为{counts[1]}个，"
        f"中等资源情景为{counts[2]}个，高资源情景为{counts[3]}个。"
        "随着所比较资源水平提高，期末构成表现为更多单元处于一类评分区间，三类单元数量减少或保持较低水平；"
        "各情景均未出现四、五类单元。该比较仅反映给定LOCF初始状态、年评分下降量、养护恢复量及资源条件下的模型模拟结果。"
    )
    discussion = (
        "等级构成图补充了长期平均评分J和最低评分未直接呈现的单元分布信息。"
        "平均评分接近的方案，各评分区间内的单元数量仍可能不同，且差异受区间阈值影响。"
        "因此，一类单元比例不能单独作为判断真实工程养护效益的依据。"
    )
    content = {"title": title, "caption": caption, "result_paragraph": result,
               "discussion_optional": discussion,
               "character_counts": {"result_paragraph": character_counts(result),
                                    "discussion_optional": character_counts(discussion)}}
    for key, low, high in [("result_paragraph", 150, 250), ("discussion_optional", 80, 120)]:
        assert all(low <= value <= high for value in content["character_counts"][key].values())
    assert not any(word in result for word in ["显著", "证明", "全局最优", "最佳预算"])
    return content


def font_run(run, size=10.5, bold=False):
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)
    fonts = run._element.get_or_add_rPr().find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        run._element.get_or_add_rPr().insert(0, fonts)
    fonts.set(qn("w:eastAsia"), "Noto Serif CJK SC")


def write_document(content, path):
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin, section.bottom_margin = Cm(2.15), Cm(2.15)
    section.left_margin, section.right_margin = Cm(2.35), Cm(2.35)
    normal = document.styles["Normal"]
    normal.font.name, normal.font.size = "Times New Roman", Pt(10.5)
    normal.paragraph_format.line_spacing = 1.2
    normal.paragraph_format.space_after = Pt(4)
    normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Noto Serif CJK SC")
    document.core_properties.title = "桥梁检测单元评分等级构成图：独立图题与论文文字"
    document.core_properties.subject = "绘图指令3；供作者审核，未插入原论文"
    document.core_properties.author = ""
    parts = [("正式图题", content["title"]), ("图注", content["caption"]),
             ("结果描述（第3章备选）", content["result_paragraph"]),
             ("讨论衔接（备选）", content["discussion_optional"])]
    for label, text in parts:
        heading = document.add_paragraph()
        heading.paragraph_format.keep_with_next = True
        heading.paragraph_format.space_before = Pt(8 if label != "正式图题" else 0)
        heading.paragraph_format.space_after = Pt(3)
        font_run(heading.add_run(label), size=11, bold=True)
        para = document.add_paragraph()
        para.paragraph_format.widow_control = True
        para.paragraph_format.keep_together = True
        para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        font_run(para.add_run(text), size=10.5, bold=(label == "正式图题"))
    document.save(path)
    # Read back every text paragraph; OOXML must preserve all mathematical symbols.
    reloaded = Document(path)
    expected = [item for part in parts for item in part]
    actual = [para.text for para in reloaded.paragraphs]
    assert actual == expected, "DOCX text read-back mismatch"
    assert len(reloaded.inline_shapes) == 0
    return len(actual)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    default_source = HERE / "input_data"
    if not default_source.is_dir():
        default_source = HERE.parent / "output_condition_composition"
    parser.add_argument("--source-dir", type=Path, default=default_source)
    parser.add_argument("--output-dir", type=Path, default=HERE)
    args = parser.parse_args()
    source_dir, out = args.source_dir.resolve(), args.output_dir.resolve()
    assert source_dir != out, "Output must be independent from the original source directory"
    out.mkdir(parents=True, exist_ok=True)
    (out / "validation").mkdir(exist_ok=True)
    observed, hashes = read_and_check(source_dir)
    content = make_content(observed)
    docx_path = out / "condition_composition_result_text.docx"
    paragraphs = write_document(content, docx_path)
    (out / "text_content.json").write_text(json.dumps(content, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    audit = {"all_text_data_checks_passed": True, "source_directory": str(source_dir),
             "source_sha256": hashes, "verified_scenarios": observed,
             "word_file": str(docx_path), "docx_paragraph_count": paragraphs,
             "readback_exact": True, "character_counts": content["character_counts"],
             "counting_method": "CJK Han characters U+4E00–U+9FFF; all non-whitespace Unicode characters, respectively.",
             "existing_manuscripts_modified": False, "source_data_modified": False}
    (out / "validation" / "text_audit.json").write_text(json.dumps(audit, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    assert all(sha256(source_dir / name) == digest for name, digest in hashes.items())
    print(json.dumps({"word": str(docx_path), "character_counts": content["character_counts"],
                      "all_text_data_checks_passed": True}, ensure_ascii=False))


if __name__ == "__main__":
    main()
